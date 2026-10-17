# 按表头对齐合并表格，保留每份输入的来源信息。 / Merge by column names and retain provenance.
"""Merge first worksheets by headers, preserve values and add source filename."""

import sys

sys.dont_write_bytecode = True
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "python-vendor"))
import datetime
import zipfile
import openpyxl
from openpyxl.cell import WriteOnlyCell
from cli_v1 import main, inputs, result, port, description_port

DESCRIPTION = {
    "protocolVersion": 1,
    "id": "table.merge_excel",
    "name": "Excel表格合并",
    "version": "1.0.0",
    "description": "Merge the first worksheet of each XLSX/XLSM by exact header; union columns and add source filename. Values only, no formatting or formula recalculation. Formula cells require saved cached results.",
    "inputs": [description_port("data"), description_port("source", "JSON", "SINGLE")],
    "outputs": [description_port("data"), description_port("source", "JSON", "SINGLE")],
    "parameters": [
        {
            "name": "sourceColumn",
            "label": "来源列名",
            "type": "string",
            "default": "来源文件",
            "required": False,
        }
    ],
}
CAPABILITY = {
    "capabilityVersion": 1,
    "toolId": DESCRIPTION["id"],
    "toolVersion": "1.0.0",
    "implementation": "AVAILABLE",
    "operation": "CONVERT",
    "inputs": [port("data", ["xlsx", "xlsm"]), port("source", ["json"], "JSON", "SINGLE")],
    "outputs": [port("data", ["xlsx"]), port("source", ["json"], "JSON", "SINGLE")],
    "knownExtensions": ["xlsx", "xlsm"],
}


def execute(request, parameters, output):
    if not set(parameters) <= {"sourceColumn"}:
        raise ValueError("Unknown merge configuration")
    source_column = parameters.get("sourceColumn", "来源文件")
    if not isinstance(source_column, str) or not source_column.strip():
        raise ValueError("来源列名不能为空")
    files = inputs(request)
    if len(files) > 500:
        raise ValueError("单次最多合并500个文件")
    columns, tables, warnings, total = [], [], [], 0
    for path, metadata in files:
        if path.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValueError("只支持XLSX/XLSM，不支持旧版XLS")
        if path.stat().st_size > 64 * 1024 * 1024:
            raise ValueError("单个工作簿超过64MiB")
        with zipfile.ZipFile(path) as archive:
            if (
                len(archive.infolist()) > 10000
                or sum(i.file_size for i in archive.infolist()) > 512 * 1024 * 1024
            ):
                raise ValueError("工作簿解压规模过大")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet, saved_sheet = workbook.worksheets[0], cached.worksheets[0]
            if (
                sheet.max_row
                and sheet.max_row > 100001
                or sheet.max_column
                and sheet.max_column > 1000
            ):
                raise ValueError("单表限制10万数据行和1000列")
            rows, saved_rows = sheet.iter_rows(), saved_sheet.iter_rows()
            header = next(rows, ())
            next(saved_rows, ())
            last = max(
                (i + 1 for i, cell in enumerate(header) if cell.value is not None), default=0
            )
            if last == 0:
                warnings.append(metadata["relativePath"] + "：空表已跳过")
                continue
            names = [str(cell.value) if cell.value is not None else "" for cell in header[:last]]
            if any(not n.strip() for n in names) or len(set(names)) != len(names):
                raise ValueError(metadata["relativePath"] + "：表头为空或重复")
            if source_column in names:
                raise ValueError("来源列名与原表头冲突，请改用其他来源列名")
            for name in names:
                if name not in columns:
                    columns.append(name)
            data = []
            for row_index, (row, saved) in enumerate(zip(rows, saved_rows), 2):
                if any(cell.value is not None for cell in row[last:]):
                    raise ValueError("存在没有表头的数据列")
                values = []
                for index, cell in enumerate(row[:last]):
                    value = cell.value
                    if cell.data_type == "f":
                        value = saved[index].value
                        if value is None:
                            raise ValueError(
                                f"{metadata['relativePath']} 第{row_index}行公式没有缓存值，请先用Excel重新计算并保存"
                            )
                    values.append(value)
                if not any(v is not None for v in values):
                    continue
                data.append(values)
                total += 1
                if total > 100000:
                    raise ValueError("合并结果超过10万行")
            tables.append((names, data, metadata["relativePath"]))
        finally:
            workbook.close()
            cached.close()
    if not columns or not tables:
        raise ValueError("没有可合并的表头")
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("合并结果")

    def append(values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)

    append(columns + [source_column])
    for names, data, name in tables:
        for row in data:
            values = dict(zip(names, row))
            append([values.get(column) for column in columns] + [name])
    filename = "合并结果.xlsx"
    (output / "payloads").mkdir()
    book.save(output / "payloads" / filename)
    reopened = openpyxl.load_workbook(output / "payloads" / filename, read_only=True)
    try:
        if next(reopened.worksheets[0].values) != tuple(columns + [source_column]):
            raise ValueError("输出表头校验失败")
    finally:
        reopened.close()
    result(request, output, filename, files, warnings)


if __name__ == "__main__":
    main(DESCRIPTION, CAPABILITY, execute)
